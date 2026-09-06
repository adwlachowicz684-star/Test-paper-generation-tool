# -*- coding: utf-8 -*-
r"""逐题核对：HTML 与 Word 两端的题干是否都在。

    python3 tools/check_render.py html/试卷_XXX.html out/试卷_XXX.docx <题目ID...>

比"统计结构数"更进一步 —— 后者只能证明数量相等，
**证明不了第 N 题在两端都存在**。数量对但内容错位的情况查不出来
（答案卷题号坑就是这么来的，见 30-pitfalls.md B4）。

做法：用题干去公式后的纯中文片段做锚点，两端各查一次。

三个注意点（都踩过）：
1. **两端公式都要置空**。HTML 的 <math>...</math> 要替换成 ''；
   Word 的公式内容在 <m:t> 里，w:t 提取出来天然为空。
   只处理一端会导致跨公式的锚点永远匹配不上。
2. **全角句号 ．(U+FF0E) ≠ 中文句号 。(U+3002)**。
   切分符漏掉它，填空题锚点会带上尾部标点而匹配失败。
3. **填空下划线 ＿＿＿ 要从锚点里去掉** —— 它是渲染出来的，不在文本里。
"""
import sys, os, re, json

sys.path.insert(0, os.path.join(
    os.path.dirname(os.path.abspath(__file__)), '..', 'py'))

import docx

BANK = 'data/bank.json'
SPLIT = r'[，。．、（）]'          # 注意含全角句号
MIN_ANCHOR = 5


def norm(t):
    return re.sub(r'\s+', '', t)


def load_texts(html_path, docx_path):
    h = open(html_path, encoding='utf-8').read()
    # HTML：公式整体置空
    htext = norm(re.sub(
        r'<[^>]+>', '',
        re.sub(r'<math\b[^>]*>.*?</math>', '', h, flags=re.S)))

    d = docx.Document(docx_path)
    allx = ''.join(p._p.xml for p in d.paragraphs)
    for t in d.tables:                      # 选项在表格里，必须算进去
        allx += t._tbl.xml
    # Word：公式内容在 m:t 里，w:t 提取出来天然为空
    wtext = norm(''.join(
        re.findall(r'<w:t[^>]*>([^<]*)</w:t>', allx)))
    return htext, wtext


def anchor_of(stem):
    # 填空位两种写法都要去：全角 ＿ 和半角 _。
    # 只处理全角时，用半角 ____ 录入的题会带着 ____ 做锚点，
    # 而 ____ 在两端都被渲染成了下划线元素（不在文本里）→ 永久匹配失败。
    s = re.sub(r'\$[^$]*\$', '', stem or '')
    s = re.sub(r'[＿_]{2,}', '', s)
    cands = [c for c in re.split(SPLIT, s) if len(norm(c)) >= MIN_ANCHOR]
    if not cands:
        return norm(s)[:10]
    return norm(max(cands, key=lambda c: len(norm(c))))


def main(argv):
    args = [a for a in argv[1:] if not a.startswith('--')]
    quiet = '--quiet' in argv            # 只打印失败项
    if len(args) < 2:
        print(__doc__)
        return 2
    html_path, docx_path = args[0], args[1]
    ids = args[2:]
    bank = json.load(open(BANK, encoding='utf-8'))
    if ids:
        sel = [q for q in bank if q['id'] in ids]
    else:
        sel = bank
    htext, wtext = load_texts(html_path, docx_path)

    ok, bad = 0, []
    for q in sel:
        a = anchor_of(q.get('stem_text'))
        inw, inh = a in wtext, a in htext
        if inw and inh:
            ok += 1
        else:
            bad.append((q['id'], a[:24], inw, inh))
    if quiet:
        # 全通过时只输出一行。题库有近 200 题，逐行打印会淹没真正的问题，
        # 也让每批核对平白多出上百行输出。
        if bad:
            for qid, a, inw, inh in bad:
                print('  ✗ %-9s %-26s W=%s H=%s'
                      % (qid, a, 'Y' if inw else 'N', 'Y' if inh else 'N'))
    else:
        for q in sel:
            a = anchor_of(q.get('stem_text'))
            inw, inh = a in wtext, a in htext
            print('  %s %-9s %-26s W=%s H=%s' % (
                '✓' if inw and inh else '✗', q['id'], a[:24],
                'Y' if inw else 'N', 'Y' if inh else 'N'))
    print('  两端一致: %d / %d' % (ok, len(sel)))
    return 0 if ok == len(sel) else 1


if __name__ == '__main__':
    sys.exit(main(sys.argv))
