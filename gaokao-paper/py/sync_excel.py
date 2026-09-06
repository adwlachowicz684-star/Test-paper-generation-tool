# -*- coding: utf-8 -*-
"""把 Excel 里人工校订过的字段同步回 bank.json

为什么需要这一步：
  拆题流程产出的 JSON 只有题干/选项/答案/解析，
  而 知识点 / 难度系数 / 题型细分 是在 Excel 里标注的。
  两端不打通的话，程序化组卷就筛不了知识点。

运行：python3 py/sync_excel.py [Excel路径]
"""
import sys, os, json, shutil

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
BANK = os.path.join(ROOT, 'data', 'bank.json')
XLSX = sys.argv[1] if len(sys.argv) > 1 else os.path.join(
    os.path.dirname(ROOT), '高三题库与间隔复习系统.xlsx')

try:
    import openpyxl
except ImportError:
    print('需要 openpyxl：pip install openpyxl')
    sys.exit(1)

if not os.path.exists(XLSX):
    print('找不到 Excel:', XLSX)
    sys.exit(1)

wb = openpyxl.load_workbook(XLSX)
ws = wb['题库索引']

# 列位（1-based）：A=1 题目ID, Y=25 难度系数, AA=27 题型细分, AB=28 知识点
C_ID, C_DIFF, C_SUBTYPE, C_KP = 1, 25, 27, 28

extra = {}
nrow = 0
for r in range(2, ws.max_row + 1):
    qid = ws.cell(row=r, column=C_ID).value
    if not qid:
        continue
    nrow += 1
    diff = ws.cell(row=r, column=C_DIFF).value
    sub = ws.cell(row=r, column=C_SUBTYPE).value
    kp = ws.cell(row=r, column=C_KP).value
    extra[str(qid).strip()] = {
        'difficulty': float(diff) if isinstance(diff, (int, float)) else None,
        'subtype': (str(sub).strip() if sub else ''),
        'kp': (str(kp).strip() if kp else ''),
    }

print(f'Excel 读取 {nrow} 行，有效 {len(extra)} 题')

bank = json.load(open(BANK, encoding='utf-8'))
hit = 0
for q in bank:
    e = extra.get(q['id'])
    if not e:
        continue
    hit += 1
    if e['difficulty'] is not None:
        q['difficulty'] = e['difficulty']
    if e['subtype']:
        q['subtype'] = e['subtype']
    if e['kp']:
        q['kp'] = e['kp']

# 原子写
tmp = BANK + '.tmp'
with open(tmp, 'w', encoding='utf-8') as f:
    json.dump(bank, f, ensure_ascii=False)
os.replace(tmp, BANK)

print(f'同步 {hit} / {len(bank)} 题')
print(f'已写入 {BANK}')

# 校验
nkp = sum(1 for q in bank if q.get('kp'))
ndf = sum(1 for q in bank if q.get('difficulty') is not None)
print(f'  含知识点 {nkp}  含难度系数 {ndf}')
